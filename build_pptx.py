# -*- coding: utf-8 -*-
"""Генерация готовой презентации проекта скейт-парка в Цхинвале."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# Палитра
BG     = RGBColor(0x0A, 0x0E, 0x1A)
BG2    = RGBColor(0x0F, 0x16, 0x28)
CARD   = RGBColor(0x14, 0x1D, 0x33)
LINE   = RGBColor(0x24, 0x31, 0x50)
ACCENT = RGBColor(0xFF, 0x57, 0x22)
GOLD   = RGBColor(0xFF, 0xC1, 0x07)
TXT    = RGBColor(0xEA, 0xF0, 0xFA)
DIM    = RGBColor(0x8B, 0x9A, 0xB8)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
blank = prs.slide_layouts[6]

FONT = "Arial"

def slide():
    s = prs.slides.add_slide(blank)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    r.fill.solid(); r.fill.fore_color.rgb = BG
    r.line.fill.background()
    r.shadow.inherit = False
    return s

def textbox(s, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tf

def setrun(r, text, size, color, bold=False, italic=False):
    r.text = text
    r.font.size = Pt(size); r.font.name = FONT
    r.font.color.rgb = color; r.font.bold = bold; r.font.italic = italic

def para(tf, text, size, color, bold=False, italic=False, space_after=6, align=PP_ALIGN.LEFT, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(space_after)
    run = p.add_run(); setrun(run, text, size, color, bold, italic)
    return p

def kicker(s, text):
    tf = textbox(s, Inches(0.85), Inches(0.7), Inches(11), Inches(0.5))
    para(tf, text.upper(), 14, GOLD, bold=True, first=True)

def card(s, l, t, w, h, icon, title, body, accent_border=False):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = CARD
    box.line.color.rgb = ACCENT if accent_border else LINE
    box.line.width = Pt(1.25 if accent_border else 1)
    box.shadow.inherit = False
    try: box.adjustments[0] = 0.06
    except Exception: pass
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.22); tf.margin_right = Inches(0.22)
    tf.margin_top = Inches(0.2); tf.margin_bottom = Inches(0.2)
    para(tf, icon, 30, GOLD, first=True, space_after=4)
    para(tf, title, 17, WHITE, bold=True, space_after=5)
    para(tf, body, 12.5, DIM, space_after=0)

def number_row(s, l, t, w, num, title, body):
    chip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, Inches(0.55), Inches(0.55))
    chip.fill.solid(); chip.fill.fore_color.rgb = ACCENT
    chip.line.fill.background(); chip.shadow.inherit = False
    ctf = chip.text_frame; ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = ctf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    setrun(p.add_run(), str(num), 20, RGBColor(0x1A,0x12,0x05), bold=True)
    tf = textbox(s, l + Inches(0.75), t - Inches(0.05), w - Inches(0.8), Inches(0.7), MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]; p.space_after = Pt(0)
    setrun(p.add_run(), title + "  ", 16, WHITE, bold=True)
    setrun(p.add_run(), body, 14, DIM)

def title_lines(s, top, lines, size=42):
    tf = textbox(s, Inches(0.85), top, Inches(11.6), Inches(1.8))
    for i, (txt, col) in enumerate(lines):
        para(tf, txt, size, col, bold=True, first=(i==0), space_after=0)

# ---------- Слайд 1: Титул ----------
s = slide()
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.25), SH)
band.fill.solid(); band.fill.fore_color.rgb = ACCENT; band.line.fill.background(); band.shadow.inherit=False
kicker(s, "Социальный проект · Южная Осетия")
tf = textbox(s, Inches(0.85), Inches(2.0), Inches(11.6), Inches(3))
para(tf, "Скейт-парк", 76, TXT, bold=True, first=True, space_after=0)
p = tf.add_paragraph(); p.space_after = Pt(0)
setrun(p.add_run(), "в ", 76, TXT, bold=True)
setrun(p.add_run(), "Цхинвале", 76, GOLD, bold=True)
tf2 = textbox(s, Inches(0.85), Inches(5.0), Inches(10), Inches(1.3))
para(tf2, "Современное и безопасное пространство для досуга и спорта\nдетей и подростков города.", 20, DIM, first=True)
pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(6.4), Inches(4.6), Inches(0.6))
pill.fill.solid(); pill.fill.fore_color.rgb = ACCENT; pill.line.fill.background(); pill.shadow.inherit=False
ptf = pill.text_frame; ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
pp = ptf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
setrun(pp.add_run(), "BMX · Самокат · Скейтборд", 18, RGBColor(0x1A,0x12,0x05), bold=True)

# ---------- Слайд 2: Проблема ----------
s = slide()
kicker(s, "Проблема")
title_lines(s, Inches(1.4), [("Детям и подросткам", TXT), ("негде проводить время", TXT)])
tf = textbox(s, Inches(0.85), Inches(3.5), Inches(11), Inches(3))
p = tf.paragraphs[0]; p.space_after = Pt(14)
setrun(p.add_run(), "В Цхинвале и в Южной Осетии в целом ", 19, TXT)
setrun(p.add_run(), "не развита инфраструктура для досуга", 19, GOLD, bold=True)
setrun(p.add_run(), " молодёжи. Современных площадок для активного отдыха почти нет.", 19, TXT)
p2 = tf.add_paragraph()
setrun(p2.add_run(), "В итоге дети катаются на оживлённых улицах, парковках и тротуарах — это ", 19, TXT)
setrun(p2.add_run(), "опасно", 19, GOLD, bold=True)
setrun(p2.add_run(), " и для них, и для прохожих. Энергия молодёжи не находит здорового выхода.", 19, TXT)

# ---------- Слайд 3: Решение ----------
s = slide()
kicker(s, "Решение")
title_lines(s, Inches(1.4), [("Построить современный", TXT), ("скейт-парк", GOLD)])
y = Inches(3.7); w = Inches(3.73); gap = Inches(0.2); x0 = Inches(0.85); h = Inches(2.9)
card(s, x0, y, w, h, "▰", "Качественное покрытие", "Рампы, фигуры и гладкое профессиональное основание.")
card(s, x0+w+gap, y, w, h, "❖", "Зона отдыха", "Скамейки, освещение, озеленение — место притяжения города.")
card(s, x0+2*(w+gap), y, w, h, "◎", "Для всех уровней", "От первых шагов новичка до тренировки сложных трюков.")

# ---------- Слайд 4: Снаряды ----------
s = slide()
kicker(s, "Спортивные направления")
title_lines(s, Inches(1.4), [("Три вида спорта на одной площадке", TXT)], size=38)
y = Inches(3.0); h = Inches(2.7)
card(s, x0, y, w, h, "BMX", "Велосипеды", "Трюковые велосипеды — сила, координация и зрелищность.")
card(s, x0+w+gap, y, w, h, "◔", "Трюковой самокат", "Самый популярный и доступный старт для детей.")
card(s, x0+2*(w+gap), y, w, h, "▱", "Скейтборд", "Классика уличного спорта: баланс и чувство тела.")
tf = textbox(s, Inches(0.85), Inches(6.1), Inches(11), Inches(0.8))
para(tf, "Дети осваивают новые спортивные снаряды и находят дело по душе.", 19, DIM, first=True)

# ---------- Слайд 5: Безопасность ----------
s = slide()
kicker(s, "Главный приоритет")
title_lines(s, Inches(1.3), [("Безопасность прежде всего", ACCENT)], size=40)
items = [
    ("Защитная экипировка.", "Шлемы, наколенники, налокотники и перчатки — обязательны для всех."),
    ("Безопасное покрытие.", "Сертифицированные материалы, продуманные радиусы фигур, без острых углов."),
    ("Зонирование.", "Раздельные зоны для новичков и опытных, чтобы потоки не пересекались."),
    ("Освещение и аптечка.", "Хороший свет вечером и медпункт первой помощи прямо на месте."),
]
yy = Inches(2.7)
for i,(t,b) in enumerate(items):
    number_row(s, Inches(0.85), yy, Inches(11.6), i+1, t, b)
    yy += Inches(1.05)

# ---------- Слайд 6: Тренер ----------
s = slide()
kicker(s, "Ключевая роль")
tf = textbox(s, Inches(0.85), Inches(1.3), Inches(11), Inches(1))
p = tf.paragraphs[0]
setrun(p.add_run(), "Нужен ", 42, TXT, bold=True)
setrun(p.add_run(), "тренер", 42, GOLD, bold=True)
tf2 = textbox(s, Inches(0.85), Inches(2.25), Inches(11.5), Inches(0.8))
para(tf2, "Без наставника спорт превращается в риск. Профессиональный тренер — основа безопасного и результативного парка.", 17, DIM, first=True)
y = Inches(3.4); w2 = Inches(5.66); h = Inches(1.75)
card(s, x0, y, w2, h, "▲", "Правильная техника", "Учит трюкам и падениям так, чтобы избежать травм.")
card(s, x0+w2+gap, y, w2, h, "◉", "Контроль и порядок", "Следит за дисциплиной, экипировкой и безопасностью на площадке.")
card(s, x0, y+h+gap, w2, h, "★", "Спортивные секции", "Группы, прогресс, соревнования и мотивация для детей.")
card(s, x0+w2+gap, y+h+gap, w2, h, "♥", "Наставничество", "Авторитетный взрослый, который направляет подростков.")

# ---------- Слайд 7: Социальная польза ----------
s = slide()
kicker(s, "Социальная сфера")
title_lines(s, Inches(1.3), [("Польза для города и его молодёжи", TXT)], size=36)
pills = ["Здоровье и спорт вместо улицы","Меньше детской преступности","Активный и полезный досуг",
         "Новые знакомства и команда","Точка притяжения города","Будущие чемпионы Осетии",
         "Гордость и развитие региона"]
px, py = Inches(0.85), Inches(2.9)
maxx = Inches(12.5); xx = px
for txt in pills:
    wpx = Inches(0.5 + len(txt)*0.115)
    if xx + wpx > maxx:
        xx = px; py += Inches(0.85)
    pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, xx, py, wpx, Inches(0.62))
    pill.fill.solid(); pill.fill.fore_color.rgb = BG2; pill.line.color.rgb = LINE
    pill.shadow.inherit=False
    ptf = pill.text_frame; ptf.vertical_anchor = MSO_ANCHOR.MIDDLE
    pp = ptf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    setrun(pp.add_run(), txt, 14, TXT, bold=True)
    xx += wpx + Inches(0.22)
tf = textbox(s, Inches(0.85), Inches(6.2), Inches(11.6), Inches(0.9))
para(tf, "Парк объединяет молодёжь вокруг здорового образа жизни и даёт детям цель.", 19, DIM, first=True)

# ---------- Слайд 8: Экономика / бюджет ----------
s = slide()
kicker(s, "Экономика и бюджет")
title_lines(s, Inches(1.3), [("Выгода для местного бизнеса и казны", TXT)], size=34)
tf = textbox(s, Inches(0.85), Inches(2.4), Inches(11.6), Inches(1.2))
p = tf.paragraphs[0]
setrun(p.add_run(), "С появлением парка вырастет спрос на спортивный инвентарь. Это ", 18, TXT)
setrun(p.add_run(), "возможность для местных торговцев", 18, GOLD, bold=True)
setrun(p.add_run(), " продавать самокаты, скейтборды, BMX, запчасти и защитную экипировку.", 18, TXT)
y = Inches(4.0); w = Inches(3.73); h = Inches(2.6)
def stat(l, big, body):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, y, w, h)
    box.fill.solid(); box.fill.fore_color.rgb = CARD; box.line.color.rgb = LINE; box.shadow.inherit=False
    tf = box.text_frame; tf.word_wrap=True
    tf.margin_left=Inches(0.25); tf.margin_top=Inches(0.3)
    para(tf, big, 34, GOLD, bold=True, first=True, space_after=8)
    para(tf, body, 15, DIM)
stat(x0, "Рост продаж", "Больше выручки у местных продавцов инвентаря и экипировки.")
stat(x0+w+gap, "Рабочие места", "Магазины, прокат, сервис, тренеры — новая занятость в городе.")
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x0+2*(w+gap), y, w, h)
box.fill.solid(); box.fill.fore_color.rgb = CARD; box.line.color.rgb = GOLD; box.line.width=Pt(1.5); box.shadow.inherit=False
tf = box.text_frame; tf.word_wrap=True; tf.margin_left=Inches(0.25); tf.margin_top=Inches(0.3)
para(tf, "В бюджет", 34, GOLD, bold=True, first=True, space_after=8)
p = tf.add_paragraph()
setrun(p.add_run(), "Налоги и сборы с торговли — ", 15, DIM)
setrun(p.add_run(), "средства в казну города", 15, WHITE, bold=True)
setrun(p.add_run(), ".", 15, DIM)

# ---------- Слайд 9: Все плюсы ----------
s = slide()
kicker(s, "Итог")
title_lines(s, Inches(1.3), [("Все плюсы проекта в одном месте", TXT)], size=36)
y = Inches(3.0); w2 = Inches(5.66); h = Inches(1.8)
card(s, x0, y, w2, h, "▱", "Спорт и развитие", "BMX, самокат, скейтборд — новые навыки для детей.")
card(s, x0+w2+gap, y, w2, h, "◈", "Безопасность", "Экипировка, зонирование, тренер и медпункт.")
card(s, x0, y+h+gap, w2, h, "♥", "Социальная польза", "Здоровый досуг и меньше уличных рисков.")
card(s, x0+w2+gap, y+h+gap, w2, h, "₽", "Доход в бюджет", "Торговля инвентарём и налоги в казну города.")

# ---------- Слайд 10: Призыв ----------
s = slide()
band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.25), SH)
band.fill.solid(); band.fill.fore_color.rgb = ACCENT; band.line.fill.background(); band.shadow.inherit=False
kicker(s, "Призыв к действию")
tf = textbox(s, Inches(0.85), Inches(1.9), Inches(11.6), Inches(2))
para(tf, "Дадим детям", 58, TXT, bold=True, first=True, space_after=0)
para(tf, "место для роста", 58, GOLD, bold=True, space_after=0)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(4.5), Inches(0.08), Inches(1.3))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background(); bar.shadow.inherit=False
tf2 = textbox(s, Inches(1.15), Inches(4.5), Inches(10), Inches(1.4), MSO_ANCHOR.MIDDLE)
para(tf2, "«Скейт-парк — это инвестиция в здоровое, активное\nи успешное будущее Южной Осетии».", 24, TXT, bold=True, italic=True, first=True)
pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(6.2), Inches(3.4), Inches(0.65))
pill.fill.solid(); pill.fill.fore_color.rgb = ACCENT; pill.line.fill.background(); pill.shadow.inherit=False
ptf = pill.text_frame; ptf.vertical_anchor=MSO_ANCHOR.MIDDLE
pp = ptf.paragraphs[0]; pp.alignment=PP_ALIGN.CENTER
setrun(pp.add_run(), "Поддержать проект", 18, RGBColor(0x1A,0x12,0x05), bold=True)
tf3 = textbox(s, Inches(4.5), Inches(6.2), Inches(4), Inches(0.65), MSO_ANCHOR.MIDDLE)
para(tf3, "Цхинвал · 2026", 18, DIM, bold=True, first=True)

prs.save("/home/user/napitki-fartuny/skatepark.pptx")
print("OK:", len(prs.slides.__iter__.__self__._sldIdLst), "слайдов")
